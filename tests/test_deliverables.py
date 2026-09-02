"""交付物中心测试：MCP 生成工具、登记/下载 API 与权限。"""
from __future__ import annotations

import base64
import tempfile
import unittest
from io import BytesIO
from pathlib import Path
from uuid import uuid4

from fastapi.testclient import TestClient
from sqlmodel import SQLModel, create_engine

import db.database as database
from config import settings
from main import app

TEST_PASSWORD = "S3ed-" + uuid4().hex[:13] + "!"


class McpGenerationToolTests(unittest.TestCase):
    def setUp(self):
        self.temp_dir = tempfile.TemporaryDirectory()
        self.workspace_root = Path(self.temp_dir.name)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_make_xlsx_creates_workbook(self):
        from mcp_server.server import make_xlsx

        reply = make_xlsx(
            "reports/交付表.xlsx",
            [["指标", "数值"], ["产量", 120], ["合格率", 0.98]],
            _workspace_root=self.workspace_root,
        )
        self.assertIn("已生成", reply)
        target = self.workspace_root / "reports" / "交付表.xlsx"
        self.assertTrue(target.exists())
        from openpyxl import load_workbook

        sheet = load_workbook(target).active
        self.assertEqual(sheet.cell(row=1, column=1).value, "指标")
        self.assertEqual(sheet.cell(row=2, column=2).value, 120)

    def test_make_docx_creates_document(self):
        from docx import Document

        from mcp_server.server import make_docx

        reply = make_docx(
            "报告.docx",
            "季度交付说明",
            ["第一段结论。", "第二段风险提示。"],
            _workspace_root=self.workspace_root,
        )
        self.assertIn("已生成", reply)
        document = Document(self.workspace_root / "报告.docx")
        texts = [paragraph.text for paragraph in document.paragraphs]
        self.assertIn("季度交付说明", texts)
        self.assertIn("第二段风险提示。", texts)

    def test_make_chart_creates_png(self):
        from mcp_server.server import make_chart

        reply = make_chart(
            "产量图",
            "bar",
            [3, 7, 5],
            ["一月", "二月", "三月"],
            "月度产量",
            _workspace_root=self.workspace_root,
        )
        self.assertIn("已生成", reply)
        target = self.workspace_root / "产量图.png"
        self.assertTrue(target.exists())
        self.assertEqual(target.read_bytes()[:8], b"\x89PNG\r\n\x1a\n")

    def test_make_chart_rejects_bad_kind_and_length(self):
        from mcp_server.server import make_chart

        with self.assertRaises(ValueError):
            make_chart("a.png", "radar", [1], ["x"], _workspace_root=self.workspace_root)
        with self.assertRaises(ValueError):
            make_chart("a.png", "bar", [1, 2], ["x"], _workspace_root=self.workspace_root)

    def test_read_file_base64_roundtrip(self):
        from mcp_server.server import read_file_base64, write_file

        write_file("hello.txt", "工作区内容", _workspace_root=self.workspace_root)
        payload = read_file_base64("hello.txt", _workspace_root=self.workspace_root)
        self.assertEqual(base64.b64decode(payload).decode("utf-8"), "工作区内容")


class DeliverableApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.temp_dir = tempfile.TemporaryDirectory()
        cls.original_upload_dir = settings.upload_dir
        cls.original_storage_backend = settings.storage_backend
        cls.original_engine = database.engine
        settings.upload_dir = str(Path(cls.temp_dir.name) / "attachments")
        settings.storage_backend = "local"
        database.engine = create_engine(
            f"sqlite:///{Path(cls.temp_dir.name, 'deliverable-test.db').as_posix()}",
            connect_args={"check_same_thread": False},
        )
        SQLModel.metadata.create_all(database.engine)
        cls.client = TestClient(app)
        cls.client.__enter__()

        owner = cls.client.post(
            "/api/v1/auth/register",
            json={
                "email": "dl-owner@example.com",
                "password": TEST_PASSWORD,
                "display_name": "DL Owner",
                "workspace_name": "DL workspace",
            },
        )
        assert owner.status_code == 201, owner.text
        cls.owner_token = owner.json()["access_token"]
        cls.workspace_id = owner.json()["workspaces"][0]["id"]

        project = cls.client.post(
            "/api/v1/projects",
            json={"name": "交付物项目"},
            headers=cls.auth_headers(),
        )
        assert project.status_code == 201, project.text
        cls.project_id = project.json()["project"]["id"]
        task = cls.client.post(
            "/api/v1/tasks",
            json={"title": "交付物样例任务", "project_id": cls.project_id},
            headers=cls.auth_headers(),
        )
        assert task.status_code == 201, task.text
        cls.task_id = task.json()["task"]["id"]

    @classmethod
    def tearDownClass(cls):
        cls.client.__exit__(None, None, None)
        database.engine.dispose()
        database.engine = cls.original_engine
        settings.upload_dir = cls.original_upload_dir
        settings.storage_backend = cls.original_storage_backend
        # Windows 上已删除交付物的文件句柄可能晚于断言释放，容忍清理竞态
        import contextlib

        with contextlib.suppress(OSError):
            cls.temp_dir.cleanup()

    @classmethod
    def auth_headers(cls, workspace_id=None):
        headers = {"Authorization": f"Bearer {cls.owner_token}"}
        if workspace_id:
            headers["X-Workspace-ID"] = workspace_id
        return headers

    def test_register_list_and_download_deliverable(self):
        from api import deliverables as deliverables_module

        png = b"\x89PNG\r\n\x1a\n" + b"fake-chart-bytes"
        original_fetch_files = deliverables_module.fetch_workspace_files
        original_fetch_bytes = deliverables_module.fetch_workspace_file_bytes

        async def fake_files(workspace_id, server="local_tools"):
            return [{"path": "chart.png", "name": "chart.png", "size": len(png)}]

        async def fake_bytes(workspace_id, path, server="local_tools"):
            if path != "chart.png":
                from fastapi import HTTPException

                raise HTTPException(status_code=404, detail="文件不存在")
            return png

        deliverables_module.fetch_workspace_files = fake_files
        deliverables_module.fetch_workspace_file_bytes = fake_bytes
        try:
            listed = self.client.get(
                "/api/v1/workspace/files",
                headers=self.auth_headers(self.workspace_id),
            )
            self.assertEqual(listed.status_code, 200, listed.text)
            self.assertEqual(listed.json()["files"][0]["path"], "chart.png")

            created = self.client.post(
                "/api/v1/deliverables",
                json={"path": "chart.png", "name": "月度产量图.png", "task_id": self.task_id},
                headers=self.auth_headers(self.workspace_id),
            )
            self.assertEqual(created.status_code, 201, created.text)
            deliverable = created.json()["deliverable"]
            self.assertEqual(deliverable["kind"], "image")
            self.assertEqual(deliverable["size_bytes"], len(png))

            empty_task = self.client.post(
                "/api/v1/deliverables",
                json={"path": "chart.png"},
                headers=self.auth_headers(self.workspace_id),
            )
            self.assertEqual(empty_task.status_code, 422)

            task_list = self.client.get(
                f"/api/v1/deliverables?task_id={self.task_id}",
                headers=self.auth_headers(self.workspace_id),
            )
            self.assertEqual(task_list.status_code, 200, task_list.text)
            self.assertEqual(len(task_list.json()["deliverables"]), 1)

            download = self.client.get(
                f"{deliverable['download_url']}",
                headers=self.auth_headers(self.workspace_id),
            )
            self.assertEqual(download.status_code, 200, download.text)
            self.assertEqual(download.content, png)

            deleted = self.client.delete(
                f"/api/v1/deliverables/{deliverable['id']}",
                headers=self.auth_headers(self.workspace_id),
            )
            self.assertEqual(deleted.status_code, 200, deleted.text)
            after = self.client.get(
                f"/api/v1/deliverables?task_id={self.task_id}",
                headers=self.auth_headers(self.workspace_id),
            )
            self.assertEqual(after.json()["deliverables"], [])
        finally:
            deliverables_module.fetch_workspace_files = original_fetch_files
            deliverables_module.fetch_workspace_file_bytes = original_fetch_bytes

    def test_deliverable_requires_task_or_conversation(self):
        response = self.client.post(
            "/api/v1/deliverables",
            json={"path": "anything.txt"},
            headers=self.auth_headers(self.workspace_id),
        )
        self.assertEqual(response.status_code, 422)


if __name__ == "__main__":
    unittest.main()
